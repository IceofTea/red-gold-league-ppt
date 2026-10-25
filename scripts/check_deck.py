# -*- coding: utf-8 -*-
"""Sanity-check a generated deck before shipping it.

Checks:
  * slide size is 16:9
  * no shape extends outside the canvas
  * every slide has text
  * pictures are embedded and non-empty
  * optional: every PPT page has a matching "第 NN 页" section in the DOCX
  * optional: warn about leftover 【placeholders】

Usage::

    python scripts/check_deck.py out/demo/demo_主题团课.pptx
    python scripts/check_deck.py out/demo/demo_主题团课.pptx --docx out/demo/demo_演讲稿.docx
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

TOL = Emu(int(0.02 * 914400))  # 0.02 inch


def check_pptx(path: Path) -> tuple[list[str], list[str], dict]:
    prs = Presentation(path)
    width, height = prs.slide_width, prs.slide_height
    errors: list[str] = []
    warnings: list[str] = []
    pictures = 0
    placeholders = 0
    media: set[str] = set()

    if abs(width / height - 16 / 9) > 0.01:
        warnings.append(
            f"slide size is {width / 914400:.2f}x{height / 914400:.2f} in, not 16:9")

    for idx, slide in enumerate(prs.slides, 1):
        has_text = False
        for sh in slide.shapes:
            if sh.left is None or sh.top is None or sh.width is None or sh.height is None:
                continue
            if (sh.left < -TOL or sh.top < -TOL
                    or sh.left + sh.width > width + TOL
                    or sh.top + sh.height > height + TOL):
                errors.append(
                    f"slide {idx}: {sh.name!r} is outside the canvas "
                    f"({sh.left / 914400:.2f},{sh.top / 914400:.2f} "
                    f"{sh.width / 914400:.2f}x{sh.height / 914400:.2f})")
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
                try:
                    blob = sh.image.blob
                    if not blob:
                        errors.append(f"slide {idx}: picture {sh.name!r} is empty")
                    else:
                        media.add(sh.image.sha1)
                except Exception as exc:  # noqa: BLE001
                    errors.append(f"slide {idx}: cannot read picture: {exc}")
            if sh.has_text_frame:
                text = sh.text_frame.text.strip()
                if text:
                    has_text = True
                    if "【" in text or "】" in text:
                        placeholders += text.count("【")
        if not has_text:
            errors.append(f"slide {idx}: no text found")

    stats = {
        "slides": len(prs.slides),
        "pictures": pictures,
        "unique_media": len(media),
        "placeholders": placeholders,
    }
    return errors, warnings, stats


def check_docx(path: Path, slides: int) -> tuple[list[str], list[str], int]:
    errors: list[str] = []
    warnings: list[str] = []
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    found = sorted({int(n) for n in re.findall(r"第\s*(\d{1,3})\s*页", text)})
    if len(found) != slides:
        errors.append(
            f"docx has {len(found)} page sections, pptx has {slides} slides")
    if found and found != list(range(1, slides + 1)):
        missing = sorted(set(range(1, slides + 1)) - set(found))
        if missing:
            errors.append(f"docx is missing pages: {missing}")
    if "【" in text:
        warnings.append("docx still contains 【placeholders】")
    return errors, warnings, len(found)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", help="generated .pptx")
    parser.add_argument("--docx", help="matching .docx script (optional)")
    args = parser.parse_args(argv)

    pptx = Path(args.pptx)
    if not pptx.exists():
        print(f"error: not found: {pptx}", file=sys.stderr)
        return 1

    errors, warnings, stats = check_pptx(pptx)
    print(f"slides:       {stats['slides']}")
    print(f"pictures:     {stats['pictures']} ({stats['unique_media']} unique)")
    print(f"placeholders: {stats['placeholders']}")

    if args.docx:
        docx = Path(args.docx)
        if not docx.exists():
            errors.append(f"docx not found: {docx}")
        else:
            d_errors, d_warnings, sections = check_docx(docx, stats["slides"])
            errors += d_errors
            warnings += d_warnings
            print(f"script pages: {sections}")

    for w in warnings:
        print(f"WARN  {w}")
    for e in errors:
        print(f"ERROR {e}")

    if errors:
        print(f"\nFAILED: {len(errors)} error(s)")
        return 1
    print("\nOK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
