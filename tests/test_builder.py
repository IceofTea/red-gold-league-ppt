# -*- coding: utf-8 -*-
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation

from red_gold_deck import build, validate_config


def make_cfg() -> dict:
    return {
        "footer": "示例团支部",
        "header": "示例团课",
        "topic": "示例主题",
        "folder": "demo",
        "ppt": "demo.pptx",
        "doc": "demo.docx",
        "expect_pages": 5,
        "pages": [
            {
                "t": "cover",
                "kicker": "★ 示例团课",
                "years": "2026",
                "title": "示例主题",
                "subtitle": "副标题",
                "lead": "第一行\n第二行",
                "big": "5",
                "script": "封面讲稿。",
            },
            {
                "t": "quote",
                "chapter": "序章",
                "small": "小标题",
                "big": "一句金句",
                "sub": "一句解释。",
                "script": "金句讲稿。",
            },
            {
                "t": "content",
                "chapter": "第一章",
                "title": "三张卡片",
                "sub": "副标题",
                "cards": [
                    ("第一", "第一张卡片的内容。"),
                    ("第二", "第二张卡片的内容。"),
                    ("第三", "第三张卡片的内容。"),
                ],
                "script": "内容页讲稿。",
            },
            {
                "t": "person",
                "chapter": "第一章",
                "title": "人物页",
                "sub": "副标题",
                "name": "【人物姓名】",
                "dates": "【生卒年】",
                "deeds": [("一", "第一件事。"), ("二", "第二件事。"), ("三", "第三件事。")],
                "script": "人物页讲稿。",
            },
            {
                "t": "thanks",
                "title": "感谢聆听",
                "sub": "示例主题",
                "script": "致谢讲稿。",
            },
        ],
    }


def test_build_minimal(tmp_path: Path) -> None:
    cfg = make_cfg()
    validate_config(cfg)
    deck = build(cfg, tmp_path)

    ppt = tmp_path / "demo" / "demo.pptx"
    doc = tmp_path / "demo" / "demo.docx"
    assert ppt.exists() and ppt.stat().st_size > 10_000
    assert doc.exists() and doc.stat().st_size > 5_000
    assert len(Presentation(ppt).slides) == 5
    assert len(deck.script) == 5

    text = "\n".join(p.text for p in Document(doc).paragraphs)
    assert "第 01 页" in text and "第 05 页" in text
    assert "示例团支部" in text


def test_unknown_page_type() -> None:
    cfg = make_cfg()
    cfg["pages"][0]["t"] = "not-a-page"
    with pytest.raises(ValueError, match="unknown page type"):
        validate_config(cfg)


def test_toc6_requires_six_items() -> None:
    cfg = make_cfg()
    cfg["pages"].append({
        "t": "toc6", "title": "目录", "sub": "sub",
        "items": [("一", "a"), ("二", "b")], "script": "s",
    })
    with pytest.raises(ValueError, match="exactly 6 items"):
        validate_config(cfg)


def test_page_count_mismatch(tmp_path: Path) -> None:
    cfg = make_cfg()
    cfg["expect_pages"] = 99
    with pytest.raises(ValueError, match="page count mismatch"):
        build(cfg, tmp_path)


def test_missing_required_key() -> None:
    cfg = make_cfg()
    del cfg["pages"][2]["cards"]
    with pytest.raises(ValueError, match="missing keys"):
        validate_config(cfg)
