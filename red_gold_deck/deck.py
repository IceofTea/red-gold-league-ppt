# -*- coding: utf-8 -*-
"""统一红金党政风 Deck：每个方法同时记录演讲稿。"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from .helpers import (
    SW, SH, RED, DEEP, BRIGHT, GOLD, LGOLD, CREAM, INK, GRAY, WHITE, PALE,
    TITLE_FONT, SERIF_FONT, BODY_FONT, NUM_FONT,
    BG_COVER, BG_PAPER, BG_CHAPTER, BG_DARK,
    add_rect, add_line, add_text, add_pic_cover, page_bg, frame, title_size,
)


class Deck:
    def __init__(self, out_path, footer_label, header_label, topic, date_label="2026年10月", key=""):
        self.out_path = out_path
        self.footer_label = footer_label
        self.header_label = header_label
        self.topic = topic
        self.date_label = date_label
        self.key = key
        self.prs = Presentation()
        self.prs.slide_width = Inches(SW)
        self.prs.slide_height = Inches(SH)
        self.n = 0
        self.script = []

    # ---------- 基础 ----------
    def _slide(self, bg=BG_PAPER):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        page_bg(s, bg)
        self.n += 1
        return s

    def say(self, title, text):
        self.script.append({"page": self.n, "title": title, "text": text.strip()})

    def header(self, s, chapter, dark=False):
        col = GOLD if dark else RED
        add_text(s, 0.78, 0.36, 7.8, 0.30,
                 [[{"text": "★  ", "size": 11, "color": col, "bold": True},
                   {"text": chapter, "size": 11, "color": LGOLD if dark else GRAY, "spacing": 1.5}]])
        add_text(s, 8.95, 0.36, 3.60, 0.30,
                 [{"text": self.header_label, "size": 9, "color": LGOLD if dark else RGBColor(0x9A, 0x8B, 0x78),
                   "align": PP_ALIGN.RIGHT}])

    def footer(self, s, dark=False, label=True):
        y = 7.08
        add_line(s, 0.78, y - 0.08, 12.55, y - 0.08,
                 LGOLD if dark else RGBColor(0xD8, 0xC8, 0xAE), 0.6, alpha=60)
        if label:
            add_text(s, 0.78, y + 0.02, 9.6, 0.26,
                     [{"text": self.footer_label, "size": 8.5,
                       "color": LGOLD if dark else RGBColor(0x9A, 0x8B, 0x78), "spacing": 0.5}])
        add_text(s, 11.65, y + 0.02, 0.90, 0.26,
                 [{"text": f"{self.n:02d}", "size": 10, "color": LGOLD if dark else RED,
                   "bold": True, "align": PP_ALIGN.RIGHT}])

    def title(self, s, text, sub=None, dark=False, y=0.82, width=11.7):
        add_text(s, 0.78, y, width, 0.66,
                 [{"text": text, "font": TITLE_FONT, "size": title_size(text),
                   "color": WHITE if dark else RED}])
        if sub:
            add_text(s, 0.80, y + 0.72, width, 0.30,
                     [{"text": sub, "size": 11.5,
                       "color": LGOLD if dark else RGBColor(0x8A, 0x7A, 0x66), "spacing": 0.3}])

    # ---------- 页面 ----------
    def cover(self, kicker, years, title, subtitle, lead, big, script):
        s = self._slide(BG_COVER)
        add_text(s, 0.92, 0.72, 9.2, 0.30,
                 [{"text": kicker, "size": 11.5, "color": LGOLD, "spacing": 1.5}])
        add_text(s, 0.92, 1.18, 9.2, 0.30,
                 [{"text": years, "size": 10, "color": GOLD, "spacing": 3}])
        n = len(title)
        csize = 62 if n <= 11 else 48 if n <= 14 else 40 if n <= 18 else 34
        add_text(s, 0.86, 1.95, 9.9, 1.25,
                 [{"text": title, "font": TITLE_FONT, "size": csize, "color": WHITE}])
        add_text(s, 0.92, 3.42, 10.6, 0.48,
                 [{"text": subtitle, "font": SERIF_FONT, "size": 21, "color": LGOLD}])
        add_line(s, 0.95, 4.22, 5.7, 4.22, GOLD, 1.1)
        add_text(s, 0.95, 4.48, 9.4, 0.80,
                 [{"text": lead, "size": 12.5, "color": RGBColor(0xE6, 0xCF, 0xB0),
                   "line_spacing": 1.65}])
        # 大数字水印
        add_text(s, 10.75, 0.45, 2.35, 1.8,
                 [{"text": big, "font": NUM_FONT, "size": 100, "color": RGBColor(0x8A, 0x3A, 0x24), "bold": True}])
        add_text(s, 0.92, 6.42, 8.4, 0.24,
                 [{"text": self.footer_label + "  ·  " + self.date_label, "size": 9, "color": LGOLD, "spacing": 1}])
        add_text(s, 9.9, 6.42, 2.6, 0.24,
                 [{"text": "THEME  ·  LECTURE", "size": 8, "color": GOLD, "spacing": 2, "align": PP_ALIGN.RIGHT}])
        self.say(title, script)

    def toc(self, title, sub, items, script,
            note="一条主线，几个问题，一次真正的思考。"):
        s = self._slide(BG_PAPER)
        self.header(s, "目录  ·  CONTENTS")
        add_text(s, 0.78, 0.82, 11.7, 0.66,
                 [{"text": title, "font": TITLE_FONT, "size": 36, "color": RED}])
        add_text(s, 0.80, 1.54, 11.7, 0.30,
                 [{"text": sub, "size": 11.5, "color": RGBColor(0x8A, 0x7A, 0x66)}])
        add_line(s, 0.80, 2.03, 12.5, 2.03, RGBColor(0xD8, 0xC8, 0xAE), 0.7, alpha=70)

        n = len(items)
        if n <= 3:
            x0, y0, cw, ch, gx, gy = 0.85, 2.45, 3.78, 2.35, 0.3, 0.2
        else:
            x0, y0, cw, ch, gx, gy = 0.85, 2.35, 5.75, 1.85, 0.32, 0.24

        for i, (head, desc) in enumerate(items):
            if n <= 3:
                x = x0 + i * (cw + gx)
                y = y0
            else:
                x = x0 + (i % 2) * (cw + gx)
                y = y0 + (i // 2) * (ch + gy)
            if n <= 3:
                add_rect(s, x - 0.12, y - 0.18, cw, ch + 0.12, PALE, alpha=55)
                add_rect(s, x - 0.12, y - 0.18, cw, 0.05, RED)
            add_text(s, x, y - 0.10, 1.2, 0.45,
                     [{"text": f"{i+1:02d}", "font": NUM_FONT, "size": 24, "color": GOLD, "bold": True}])
            add_text(s, x + 0.82, y - 0.04, cw - 0.82, 0.40,
                     [{"text": head, "font": SERIF_FONT, "size": 19, "color": RED}])
            add_text(s, x + 0.82, y + 0.44, cw - 0.95, ch - 0.55,
                     [{"text": desc, "size": 10.5, "color": GRAY, "line_spacing": 1.55}])
            add_line(s, x + 0.82, y + ch - 0.10, x + cw - 0.1, y + ch - 0.10,
                     RGBColor(0xD8, 0xC8, 0xAE), 0.6, alpha=55)
        if note:
            add_text(s, 0.85, 6.15, 11.7, 0.40,
                     [{"text": note, "size": 14,
                       "color": RGBColor(0x8A, 0x4A, 0x2A), "italic": True}])
        self.footer(s)
        self.say(title, script)

    def chapter(self, num, cn, title, quote, script, img=None, img_caption=""):
        s = self._slide(BG_CHAPTER)
        if img:
            add_pic_cover(s, img, 8.52, 1.25, 3.98, 5.15, crop_bias_y=0.42)
            frame(s, 8.52, 1.25, 3.98, 5.15, GOLD, 1.0)
            add_rect(s, 8.52, 1.25, 3.98, 5.15, DEEP, alpha=24)
            if img_caption:
                add_text(s, 8.55, 6.48, 3.9, 0.24,
                         [{"text": img_caption, "size": 8, "color": GOLD, "spacing": 2}])
        add_text(s, 0.62, 0.40, 4.2, 1.9,
                 [{"text": num, "font": NUM_FONT, "size": 118, "color": LGOLD, "bold": True}])
        add_text(s, 0.78, 2.42, 6.2, 0.30,
                 [{"text": f"第{cn}章  /  CHAPTER {num}", "size": 11.5, "color": LGOLD, "spacing": 3}])
        tn = len(title)
        csize = 48 if tn <= 10 else 42 if tn <= 12 else 36 if tn <= 14 else 32 if tn <= 16 else 28
        add_text(s, 0.78, 2.82, 7.35, 1.0,
                 [{"text": title, "font": TITLE_FONT, "size": csize, "color": WHITE}])
        add_line(s, 0.80, 4.05, 4.55, 4.05, GOLD, 1.2)
        add_text(s, 0.80, 4.28, 7.15, 1.5,
                 [{"text": quote, "size": 16, "color": LGOLD, "line_spacing": 1.75}])
        self.footer(s, dark=True)
        self.say(title, script)

    def content(self, chapter, title, sub, cards, script, lead=None, cols=3, note=None):
        s = self._slide(BG_PAPER)
        self.header(s, chapter)
        self.title(s, title, sub)
        y = 2.62
        if lead:
            add_text(s, 0.85, 1.97, 11.7, 0.55,
                     [{"text": lead, "size": 12, "color": INK, "line_spacing": 1.55}])
            y = 2.72
        n = len(cards)
        if n <= 3:
            w, gap, x0 = 3.75, 0.30, 0.85
        else:
            w, gap, x0 = 2.85, 0.20, 0.85
        h = 2.90 if lead else 3.10
        hsize = 16.5 if n <= 3 else 15
        for i, (head, body) in enumerate(cards):
            x = x0 + i * (w + gap)
            add_rect(s, x, y, w, h, PALE, alpha=60)
            add_rect(s, x, y, w, 0.055, RED)
            add_text(s, x + 0.22, y + 0.18, 0.8, 0.35,
                     [{"text": f"{i+1:02d}", "font": NUM_FONT, "size": 17, "color": GOLD, "bold": True}])
            add_text(s, x + 0.22, y + 0.60, w - 0.44, 0.52,
                     [{"text": head, "font": SERIF_FONT, "size": hsize, "color": RED, "line_spacing": 1.2}])
            add_text(s, x + 0.22, y + 1.15, w - 0.42, h - 1.28,
                     [{"text": body, "size": 11, "color": GRAY, "line_spacing": 1.7}])
        if note:
            add_text(s, 0.85, min(y + h + 0.18, 6.42), 11.7, 0.32,
                     [{"text": note, "size": 10.8, "color": RGBColor(0x8A, 0x4A, 0x2A),
                       "italic": True}])
        self.footer(s)
        self.say(title, script)

    def timeline(self, chapter, title, sub, events, script, note=None):
        s = self._slide(BG_PAPER)
        self.header(s, chapter)
        self.title(s, title, sub)
        yline = 3.62
        x_start, x_end = 1.55, 11.78
        add_line(s, x_start, yline, x_end, yline, RGBColor(0xC8, 0xA4, 0x5D), 1.1)
        n = len(events)
        step = (x_end - x_start) / max(n - 1, 1)
        for i, (year, head, desc) in enumerate(events):
            x = x_start + i * step
            add_rect(s, x - 0.07, yline - 0.07, 0.14, 0.14, RED, line=GOLD, lw=0.8)
            add_text(s, x - 1.00, yline - 0.75, 2.0, 0.40,
                     [{"text": year, "font": NUM_FONT, "size": 17, "color": RED, "bold": True,
                       "align": PP_ALIGN.CENTER}])
            add_text(s, x - 1.05, yline + 0.18, 2.1, 0.42,
                     [{"text": head, "size": 12.5, "color": INK, "bold": True,
                       "align": PP_ALIGN.CENTER}])
            add_text(s, x - 1.08, yline + 0.66, 2.16, 1.25,
                     [{"text": desc, "size": 9.2, "color": GRAY, "line_spacing": 1.45,
                       "align": PP_ALIGN.CENTER}])
        if note:
            add_text(s, 0.85, 6.22, 11.7, 0.55,
                     [{"text": note, "size": 11.2, "color": RGBColor(0x8A, 0x4A, 0x2A),
                       "italic": True, "line_spacing": 1.4}])
        self.footer(s)
        self.say(title, script)

    def stats(self, chapter, title, sub, numbers, script, note=None):
        s = self._slide(BG_PAPER)
        self.header(s, chapter)
        self.title(s, title, sub)
        n = len(numbers)
        w = 11.55 / n
        for i, (num, label, desc) in enumerate(numbers):
            x = 0.85 + i * w
            add_text(s, x, 2.55, w - 0.2, 0.95,
                     [{"text": num, "font": NUM_FONT, "size": 50, "color": RED, "bold": True}])
            add_text(s, x, 3.55, w - 0.2, 0.38,
                     [{"text": label, "font": SERIF_FONT, "size": 15, "color": INK}])
            add_text(s, x, 4.05, w - 0.28, 1.35,
                     [{"text": desc, "size": 10.3, "color": GRAY, "line_spacing": 1.55}])
            if i < n - 1:
                add_line(s, x + w - 0.12, 2.65, x + w - 0.12, 5.55,
                         RGBColor(0xD8, 0xC8, 0xAE), 0.7, alpha=60)
        if note:
            add_line(s, 0.85, 6.05, 12.45, 6.05, RGBColor(0xD8, 0xC8, 0xAE), 0.7, alpha=60)
            add_text(s, 0.85, 6.20, 11.7, 0.55,
                     [{"text": note, "size": 11.5, "color": RGBColor(0x8A, 0x4A, 0x2A),
                       "italic": True, "line_spacing": 1.4}])
        self.footer(s)
        self.say(title, script)

    def quote(self, chapter, small, big, sub, script, note=None):
        s = self._slide(BG_CHAPTER)
        self.header(s, chapter, dark=True)
        add_text(s, 0.88, 1.22, 10.0, 0.34,
                 [{"text": small, "size": 11, "color": GOLD, "spacing": 3}])
        bn = len(big)
        qsize = 48 if bn <= 17 else 40 if bn <= 24 else 34 if bn <= 32 else 28
        add_text(s, 0.88, 1.78, 11.5, 2.1,
                 [{"text": big, "font": TITLE_FONT, "size": qsize,
                   "color": WHITE, "line_spacing": 1.35}])
        add_line(s, 0.92, 4.18, 5.35, 4.18, GOLD, 1.2)
        add_text(s, 0.92, 4.42, 10.8, 0.95,
                 [{"text": sub, "size": 16, "color": LGOLD, "line_spacing": 1.65}])
        if note:
            add_text(s, 0.92, 5.85, 10.8, 0.55,
                     [{"text": note, "size": 10, "color": RGBColor(0xC8, 0xA4, 0x5D), "italic": True}])
        self.footer(s, dark=True)
        self.say(big, script)

    def two_col(self, chapter, title, sub, left_title, left_items,
                right_title, right_items, script, note=None):
        s = self._slide(BG_PAPER)
        self.header(s, chapter)
        self.title(s, title, sub)
        for side, (x, head, items, color) in enumerate([
            (0.85, left_title, left_items, RED),
            (6.90, right_title, right_items, DEEP),
        ]):
            add_rect(s, x, 2.12, 5.55, 0.46, color)
            add_text(s, x + 0.20, 2.18, 5.15, 0.30,
                     [{"text": head, "size": 12, "color": WHITE, "bold": True}])
            yy = 2.78
            for h, b in items:
                add_text(s, x + 0.12, yy, 0.16, 0.28,
                         [{"text": "◆", "size": 8, "color": GOLD}])
                add_text(s, x + 0.38, yy - 0.02, 5.0, 0.30,
                         [{"text": h, "size": 12, "color": color, "bold": True}])
                add_text(s, x + 0.38, yy + 0.31, 4.95, 0.55,
                         [{"text": b, "size": 9.6, "color": GRAY, "line_spacing": 1.45}])
                yy += 0.92
        if note:
            add_text(s, 0.85, 6.48, 11.7, 0.30,
                     [{"text": note, "size": 10.5, "color": RGBColor(0x8A, 0x4A, 0x2A), "italic": True}])
        self.footer(s)
        self.say(title, script)

    def closing(self, kicker, title, lines, slogan, script, note=None):
        s = self._slide(BG_COVER)
        add_text(s, 0.95, 1.30, 10.0, 0.30,
                 [{"text": kicker, "size": 11, "color": GOLD, "spacing": 3}])
        add_text(s, 0.90, 1.85, 11.2, 0.95,
                 [{"text": title, "font": TITLE_FONT, "size": title_size(title, 48), "color": WHITE}])
        yy = 3.10
        for line in lines:
            add_text(s, 0.95, yy, 11.0, 0.42,
                     [{"text": line, "size": 16.5, "color": RGBColor(0xE6, 0xCF, 0xB0)}])
            yy += 0.56
        add_line(s, 0.98, 5.65, 5.4, 5.65, GOLD, 1.1)
        add_text(s, 0.98, 5.88, 11.2, 0.62,
                 [{"text": slogan, "font": SERIF_FONT, "size": 23, "color": LGOLD}])
        if note:
            add_text(s, 0.95, 6.62, 9.5, 0.24,
                     [{"text": note, "size": 8.8, "color": RGBColor(0xB9, 0x93, 0x62)}])
        self.footer(s, dark=True, label=False)
        self.say(title, script)

    def sources(self, items, script, title="图片与资料说明",
                sub="感谢所有记录、保存和传播历史的人",
                footer_note="主题团课学习资料"):
        s = self._slide(BG_PAPER)
        self.header(s, "附录  ·  资料说明")
        self.title(s, title, sub)
        n = len(items)
        w = 11.6 / n
        body_size = 9.5 if n <= 5 else 8.6
        head_size = 17 if n <= 5 else 15
        for i, (head, body) in enumerate(items):
            x = 0.85 + i * w
            add_text(s, x, 2.12, 0.9, 0.35,
                     [{"text": f"{i+1:02d}", "font": NUM_FONT, "size": 15, "color": GOLD, "bold": True}])
            add_text(s, x, 2.58, w - 0.3, 0.40,
                     [{"text": head, "font": SERIF_FONT, "size": head_size, "color": RED}])
            add_text(s, x, 3.12, w - 0.32, 2.7,
                     [{"text": body, "size": body_size, "color": GRAY, "line_spacing": 1.55}])
            if i < n - 1:
                add_line(s, x + w - 0.15, 2.15, x + w - 0.15, 5.9,
                         RGBColor(0xD8, 0xC8, 0xAE), 0.7, alpha=60)
        add_text(s, 0.85, 6.25, 11.7, 0.30,
                 [{"text": self.footer_label + "  ·  " + footer_note, "size": 9.5,
                   "color": RGBColor(0x9A, 0x8B, 0x78), "spacing": 1}])
        self.footer(s)
        self.say(title, script)

    def save(self):
        out_dir = os.path.dirname(os.path.abspath(self.out_path))
        os.makedirs(out_dir, exist_ok=True)
        self.prs.save(self.out_path)
        return self.out_path
