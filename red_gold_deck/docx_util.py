# -*- coding: utf-8 -*-
"""把 Deck.script 生成可直接照读的 Word 演讲稿。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _set_font(run, name="宋体", size=12, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def _footer(section, text):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    _set_font(run, "宋体", 9, color=RGBColor(0x88, 0x88, 0x88))


def build_docx(deck, path, title, subtitle, footer,
               note="本稿与 PPT 逐页对应，可直接照读。方括号【】内为提示，不读。"):
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)

    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    # 封面标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    p.paragraph_format.space_after = Pt(10)
    _set_font(p.add_run(title), "方正小标宋简体", 24, bold=True, color=RGBColor(0x9E, 0x1B, 0x1B))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    _set_font(p.add_run(subtitle), "华文中宋", 14, color=RGBColor(0x6E, 0x0C, 0x0C))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    _set_font(p.add_run(footer), "宋体", 10.5, color=RGBColor(0x66, 0x66, 0x66))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    _set_font(p.add_run(note), "宋体", 9.5, color=RGBColor(0x88, 0x88, 0x88))

    # 正文
    total = len(deck.script)
    for item in deck.script:
        page = item["page"]
        head = item["title"]

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        _set_font(p.add_run(f"第 {page:02d} 页  ｜  {head}"), "黑体", 13, bold=True, color=RGBColor(0x9E, 0x1B, 0x1B))

        for para in item["text"].split("\n"):
            para = para.strip()
            if not para:
                continue
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(24)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            p.paragraph_format.space_after = Pt(2)
            _set_font(p.add_run(para), "宋体", 12)

    # 结束语
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    _set_font(p.add_run("谢谢大家！"), "华文中宋", 15, bold=True, color=RGBColor(0x9E, 0x1B, 0x1B))

    _footer(section, f"{title}  ·  主题团课演讲稿  ·  共 {total} 页")
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    doc.save(path)
    return path
