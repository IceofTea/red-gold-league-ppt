# -*- coding: utf-8 -*-
"""Shared layout helpers for the red-gold Chinese lecture deck.

The module is deliberately dependency-light: python-pptx, Pillow and lxml.
All paths are resolved relative to the installed package, so the deck can be
generated from any working directory.
"""
from __future__ import annotations

import os
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SW, SH = 13.333, 7.5

RED = RGBColor(0x9E, 0x1B, 0x1B)
DEEP = RGBColor(0x6E, 0x0C, 0x0C)
BRIGHT = RGBColor(0xC0, 0x27, 0x1E)
GOLD = RGBColor(0xC8, 0xA4, 0x5D)
LGOLD = RGBColor(0xE8, 0xD3, 0xA3)
CREAM = RGBColor(0xF7, 0xF1, 0xE5)
INK = RGBColor(0x2B, 0x24, 0x1E)
GRAY = RGBColor(0x6F, 0x65, 0x5B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xFB, 0xF6, 0xEC)

# Fonts are intentionally configurable: the default stack targets a Windows
# machine with Chinese government-document fonts installed.  On macOS/Linux,
# set the four environment variables to any installed CJK family.
TITLE_FONT = os.environ.get("RED_GOLD_TITLE_FONT", "方正小标宋简体")
SERIF_FONT = os.environ.get("RED_GOLD_SERIF_FONT", "华文中宋")
BODY_FONT = os.environ.get("RED_GOLD_BODY_FONT", "微软雅黑")
NUM_FONT = os.environ.get("RED_GOLD_NUM_FONT", "Noto Serif SC")

ASSETS_DIR = Path(os.environ.get(
    "RED_GOLD_ASSETS", Path(__file__).resolve().parent / "assets"))
BG_DIR = ASSETS_DIR / "bg"

BG_COVER = str(BG_DIR / "cover_bg.jpg")
BG_PAPER = str(BG_DIR / "paper_bg.jpg")
BG_CHAPTER = str(BG_DIR / "chapter_bg.jpg")
BG_DARK = str(BG_DIR / "dark_bg.jpg")


def _ea(run, name: str) -> None:
    """Set the East-Asian and complex-script typeface on a run."""
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", name)


def no_shadow(shape) -> None:
    spPr = shape._element.spPr
    for e in spPr.findall(qn("a:effectLst")):
        spPr.remove(e)
    etree.SubElement(spPr, qn("a:effectLst"))


def add_rect(slide, x, y, w, h, color, alpha=None, line=None, lw=0.75,
             shape=MSO_SHAPE.RECTANGLE, radius=None):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
    if alpha is not None:
        srgb = shp.fill.fore_color._xFill.find(qn("a:srgbClr"))
        a = etree.SubElement(srgb, qn("a:alpha"))
        a.set("val", str(int(alpha * 1000)))
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    no_shadow(shp)
    return shp


def frame(slide, x, y, w, h, color=GOLD, lw=0.9):
    add_rect(slide, x, y, w, h, RGBColor(0, 0, 0), alpha=0, line=color, lw=lw)


def add_line(slide, x1, y1, x2, y2, color, width=1.0, dash=None, alpha=None):
    from pptx.enum.shapes import MSO_CONNECTOR
    ln = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(width)
    if dash:
        ln.line.dash_style = dash
    if alpha is not None:
        srgb = ln.line.color._xFill.find(qn("a:srgbClr"))
        a = etree.SubElement(srgb, qn("a:alpha"))
        a.set("val", str(int(alpha * 1000)))
    return ln


def add_text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             line_spacing=1.25, space_after=0, margins=(0, 0, 0, 0), wrap=True):
    """Add a text box.

    ``runs`` is a list of paragraphs.  A paragraph may be a dict, or a list of
    dicts for mixed formatting within one line.
    """
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom = [
        Inches(m) for m in margins]
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if isinstance(para, dict):
            para = [para]
        p.alignment = para[0].get("align", align)
        p.line_spacing = para[0].get("line_spacing", line_spacing)
        p.space_before = Pt(para[0].get("space_before", 0))
        p.space_after = Pt(para[0].get("space_after", space_after))
        for rd in para:
            r = p.add_run()
            r.text = rd.get("text", "")
            fname = rd.get("font", BODY_FONT)
            r.font.name = fname
            r.font.size = Pt(rd.get("size", 18))
            r.font.bold = rd.get("bold", False)
            r.font.italic = rd.get("italic", False)
            r.font.color.rgb = rd.get("color", INK)
            r.font.language_id = MSO_LANGUAGE_ID.SIMPLIFIED_CHINESE
            _ea(r, fname)
            if rd.get("spacing"):
                r.font._rPr.set("spc", str(int(rd["spacing"] * 100)))
    return tb


def add_pic_cover(slide, path, x, y, w, h, crop_bias_x=0.5, crop_bias_y=0.5):
    """Insert a picture and crop it to fill the target rectangle."""
    iw, ih = Image.open(path).size
    target = w / h
    src = iw / ih
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                   Inches(w), Inches(h))
    if src > target:
        frac = 1 - target / src
        pic.crop_left = frac * crop_bias_x
        pic.crop_right = frac * (1 - crop_bias_x)
    elif src < target:
        frac = 1 - src / target
        pic.crop_top = frac * crop_bias_y
        pic.crop_bottom = frac * (1 - crop_bias_y)
    return pic


def pic_alpha(pic, pct):
    blip = pic._element.blipFill.blip
    amf = etree.SubElement(blip, qn("a:alphaModFix"))
    amf.set("amt", str(int(pct * 1000)))


def page_bg(slide, path):
    return slide.shapes.add_picture(str(path), 0, 0, Inches(SW), Inches(SH))


def title_size(text: str, base: int = 38) -> int:
    n = len(text)
    if n <= 18:
        return base
    if n <= 22:
        return base - 4
    if n <= 26:
        return base - 7
    return base - 10
